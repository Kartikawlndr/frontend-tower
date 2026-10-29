#!/usr/bin/env python3
"""Generate the bulk tower-registration workbook template."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = (
    ROOT
    / "extensions"
    / "endpoints"
    / "bulk-tower-upload"
    / "templates"
    / "Template_Registrasi_Menara.xlsx"
)
DATA_ROWS = 1000
FIRST_DATA_ROW = 5

FORM_HEADERS = (
    "Kode/ID Menara",
    "Nama Pemilik Menara Telekomunikasi",
    "Tahun Pembangunan Menara Telekomunikasi",
    "Penyelenggara Jaringan Telekomunikasi",
    "Longitude",
    "Latitude",
    "Alamat",
    "Jenis Menara Telekomunikasi",
    "Tinggi Menara Telekomunikasi (AGL)",
    "Kapasitas maksimal menara telekomunikasi",
    "Kapasitas menara telekomunikasi yang disewakan",
    "Kapasitas menara telekomunikasi yang digunakan sendiri",
)
FORM_NOTES = (
    "Wajib; teks maks. 50 karakter. Masukkan kode mentah sebagai teks, contoh 000101.",
    "Wajib; teks maks. 50 karakter.",
    "Wajib; pilih tahun dari 1900 sampai tahun berjalan.",
    "Wajib; teks maks. 50 karakter. Pisahkan nama multi-tenant dengan koma.",
    "Wajib; angka 95 sampai 141, maksimal 6 angka desimal.",
    "Wajib; angka -11 sampai 6, maksimal 6 angka desimal.",
    "Wajib; teks maks. 50 karakter. Masukkan detail jalan/RT/RW.",
    "Wajib; teks maks. 50 karakter.",
    "Wajib; bilangan bulat 1 sampai 999.",
    "Wajib; teks maks. 50 karakter.",
    "Wajib; teks maks. 50 karakter.",
    "Wajib; teks maks. 50 karakter.",
)

NAVY = "17365D"
BLUE = "DDEBF7"
BLUE_DARK = "1F4E78"
GREEN_DARK = "38761D"
WHITE = "FFFFFF"
BLACK = "1F1F1F"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
THIN_BLUE = "B4C7E7"


def safe_name(code: str, prefix: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_]", "_", code)
    if not value or value[0].isdigit():
        value = f"_{value}"
    return f"list_{prefix}_{value}"


def add_defined_name(workbook: Workbook, name: str, reference: str) -> None:
    workbook.defined_names.add(DefinedName(name, attr_text=reference))


def quoted_range(sheet_name: str, column: str, first: int, last: int) -> str:
    return f"'{sheet_name}'!${column}${first}:${column}${last}"


def clean_display(value: Any) -> str:
    return str(value).replace("_", " ").strip()


def load_regions(path: Path) -> list[dict[str, Any]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    regions = data.get("regions")
    if not isinstance(regions, list) or len(regions) != 38:
        raise ValueError(f"Data wilayah harus berisi 38 provinsi: {path}")
    return regions


def style_header(cell, fill_color: str, font_color: str = WHITE) -> None:
    cell.fill = PatternFill("solid", fgColor=fill_color)
    cell.font = Font(name="Aptos", bold=True, color=font_color)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style="thin", color=WHITE),
        right=Side(style="thin", color=WHITE),
        top=Side(style="thin", color=WHITE),
        bottom=Side(style="thin", color=WHITE),
    )


def set_auto_width(sheet: Worksheet, minimum: int = 12, maximum: int = 48) -> None:
    for column_cells in sheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        longest = 0
        for cell in column_cells[: min(sheet.max_row, 250)]:
            if cell.value is not None:
                longest = max(longest, len(str(cell.value)))
        sheet.column_dimensions[column_letter].width = max(
            minimum, min(maximum, longest + 2)
        )


def add_list_validation(
    sheet: Worksheet, formula: str, cell_range: str, prompt: str
) -> None:
    validation = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=False,
        errorStyle="stop",
        showErrorMessage=True,
        errorTitle="Pilihan tidak valid",
        error=prompt,
        showInputMessage=True,
        promptTitle="Pilih dari daftar",
        prompt=prompt,
    )
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def add_custom_validation(
    sheet: Worksheet, formula: str, cell_range: str, title: str, message: str
) -> None:
    validation = DataValidation(
        type="custom",
        formula1=formula,
        allow_blank=False,
        errorStyle="stop",
        showErrorMessage=True,
        errorTitle=title,
        error=message,
    )
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def create_data_master(
    workbook: Workbook, regions: list[dict[str, Any]], regions_path: Path
) -> tuple[Worksheet, dict[str, int]]:
    sheet = workbook.create_sheet("Data_Master")
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False

    headers = (
        "Provinsi",
        "Provinsi Pencarian",
        "Daftar Kabupaten/Kota",
        "Kabupaten/Kota",
        "Relasi Provinsi dan Kabupaten/Kota",
        "Daftar Kecamatan",
        "Kecamatan",
        "Relasi Kabupaten/Kota dan Kecamatan",
        "Daftar Kelurahan",
        "Kelurahan",
        "Relasi Kecamatan dan Kelurahan",
        "Kode Wilayah Kelurahan",
        "Tahun Pembangunan",
    )
    for column, header in enumerate(headers, 1):
        style_header(sheet.cell(1, column, header), BLUE_DARK)

    province_rows: list[tuple[str, str, str]] = []
    regency_rows: list[tuple[str, str, str]] = []
    district_rows: list[tuple[str, str, str]] = []
    village_rows: list[tuple[str, str, str]] = []
    province_ranges: list[tuple[str, list[str]]] = []
    regency_ranges: list[tuple[str, list[str]]] = []
    district_ranges: list[tuple[str, list[str]]] = []
    stats = {"provinces": 0, "regencies": 0, "districts": 0, "villages": 0}

    for province in regions:
        province_name = clean_display(province["name"])
        province_code = str(province["code"])
        province_list_name = safe_name(province_code, "prov")
        province_regencies = province.get("kabKota", [])
        province_ranges.append(
            (
                province_list_name,
                [clean_display(regency["name"]) for regency in province_regencies],
            )
        )
        province_rows.append((province_name, province_name, province_list_name))
        stats["provinces"] += 1

        for regency in province_regencies:
            regency_name = clean_display(regency["name"])
            regency_code = str(regency["code"])
            regency_list_name = safe_name(regency_code, "kab")
            regency_parent = f"{province_name}|{regency_name}"
            regency_districts = regency.get("kecamatan", [])
            regency_ranges.append(
                (
                    regency_list_name,
                    [clean_display(district["name"]) for district in regency_districts],
                )
            )
            regency_rows.append(
                (regency_name, regency_parent, regency_list_name)
            )
            stats["regencies"] += 1

            for district in regency_districts:
                district_name = clean_display(district["name"])
                district_code = str(district["code"])
                district_list_name = safe_name(district_code, "kec")
                district_parent = f"{province_name}|{regency_name}|{district_name}"
                villages = district.get("kelurahan", [])
                district_ranges.append(
                    (
                        district_list_name,
                        [clean_display(village["name"]) for village in villages],
                    )
                )
                district_rows.append(
                    (district_name, district_parent, district_list_name)
                )
                stats["districts"] += 1

                for village in villages:
                    village_name = clean_display(village["name"])
                    village_code = str(village["code"])
                    village_rows.append(
                        (
                            village_name,
                            district_parent,
                            village_code,
                        )
                    )
                    stats["villages"] += 1

    province_start = 2
    for row_index, row in enumerate(province_rows, province_start):
        for column, value in enumerate(row, 1):
            sheet.cell(row_index, column, value)
    province_end = province_start + len(province_rows) - 1

    regency_start = 2
    for row_index, row in enumerate(regency_rows, regency_start):
        for column, value in enumerate(row, 4):
            sheet.cell(row_index, column, value)
    regency_end = regency_start + len(regency_rows) - 1

    district_start = 2
    for row_index, row in enumerate(district_rows, district_start):
        for column, value in enumerate(row, 7):
            sheet.cell(row_index, column, value)
    district_end = district_start + len(district_rows) - 1

    village_start = 2
    for row_index, row in enumerate(village_rows, village_start):
        for column, value in enumerate(row, 10):
            sheet.cell(row_index, column, value)
    village_end = village_start + len(village_rows) - 1

    for row_index, year in enumerate(range(1900, 10000), 2):
        sheet.cell(row_index, 13, year)

    add_defined_name(
        workbook,
        "province_list",
        quoted_range("Data_Master", "A", province_start, province_end),
    )
    add_defined_name(
        workbook,
        "province_lookup_names",
        quoted_range("Data_Master", "B", province_start, province_end),
    )
    add_defined_name(
        workbook,
        "province_range_keys",
        quoted_range("Data_Master", "C", province_start, province_end),
    )
    add_defined_name(
        workbook,
        "regency_lookup_keys",
        quoted_range("Data_Master", "E", regency_start, regency_end),
    )
    add_defined_name(
        workbook,
        "regency_range_keys",
        quoted_range("Data_Master", "F", regency_start, regency_end),
    )
    add_defined_name(
        workbook,
        "district_lookup_keys",
        quoted_range("Data_Master", "H", district_start, district_end),
    )
    add_defined_name(
        workbook,
        "district_range_keys",
        quoted_range("Data_Master", "I", district_start, district_end),
    )
    add_defined_name(
        workbook,
        "year_list",
        "'Data_Master'!$M$2:INDEX('Data_Master'!$M:$M,YEAR(TODAY())-1898)",
    )

    for name, values in (
        ("province_ranges", province_ranges),
        ("regency_ranges", regency_ranges),
        ("district_ranges", district_ranges),
    ):
        column = {"province_ranges": "D", "regency_ranges": "G", "district_ranges": "J"}[name]
        start = {
            "province_ranges": province_start,
            "regency_ranges": regency_start,
            "district_ranges": district_start,
        }[name]
        for list_name, entries in values:
            first = start
            last = first + len(entries) - 1
            if entries:
                add_defined_name(
                    workbook,
                    list_name,
                    quoted_range("Data_Master", column, first, last),
                )
            start = last + 1

    source = json.loads(regions_path.read_text(encoding="utf-8"))
    sheet["O1"] = "Sumber data"
    sheet["O1"].font = Font(bold=True, color=WHITE)
    sheet["O1"].fill = PatternFill("solid", fgColor=BLUE_DARK)
    sheet["O2"] = source.get("source", "")
    sheet["O3"] = source.get("dataReference", "")
    sheet["O4"] = f"Diperbarui: {source.get('updatedAt', 'tidak diketahui')}"
    sheet["O5"] = source.get("attribution", "")
    sheet.column_dimensions["O"].width = 55

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=13):
        for cell in row:
            cell.alignment = Alignment(vertical="top")
    for column in range(1, 14):
        sheet.column_dimensions[get_column_letter(column)].width = 34
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["M"].width = 20
    return sheet, stats


def create_form_sheet(workbook: Workbook) -> Worksheet:
    sheet = workbook.active
    sheet.title = "Form_Registrasi"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:4"

    sheet.merge_cells("A1:L1")
    title = sheet["A1"]
    title.value = "TEMPLATE UPLOAD BULK REGISTRASI MENARA TELEKOMUNIKASI"
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.font = Font(name="Aptos Display", size=16, bold=True, color=WHITE)
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 32

    sheet.merge_cells("A2:L2")
    instructions = sheet["A2"]
    instructions.value = (
        "Isi satu baris untuk satu menara. Kolom bertanda Wajib harus diisi. "
        "Provinsi, Kabupaten/Kota, Kecamatan, dan Kelurahan diisi otomatis dari koordinat. "
        "Titik di laut yang tidak masuk poligon wilayah akan memiliki data wilayah kosong. "
        "Kode/ID menara harus berupa teks agar angka nol di depan tetap tersimpan. "
        "Maksimal 1.000 baris data."
    )
    instructions.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    instructions.font = Font(name="Aptos", size=10, color=BLACK)
    instructions.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 34

    for column, (header, note) in enumerate(zip(FORM_HEADERS, FORM_NOTES), 1):
        style_header(sheet.cell(3, column, header), BLUE_DARK)
        sheet.cell(4, column, note)
        sheet.cell(4, column).font = Font(name="Aptos", size=9, italic=True, color=GRAY)
        sheet.cell(4, column).fill = PatternFill("solid", fgColor="EAF2F8")
        sheet.cell(4, column).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        sheet.cell(4, column).border = Border(
            bottom=Side(style="thin", color=THIN_BLUE)
        )

    sheet.row_dimensions[3].height = 62
    sheet.row_dimensions[4].height = 54
    final_row = FIRST_DATA_ROW + DATA_ROWS - 1
    for row in range(FIRST_DATA_ROW, final_row + 1):
        sheet.row_dimensions[row].height = 22
        for column in range(1, 13):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Aptos", size=10, color=BLACK)
            cell.alignment = Alignment(vertical="center")
            if column == 3:
                cell.fill = PatternFill("solid", fgColor=BLUE)
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            if column in (1, 2, 4, 7, 8, 10, 11, 12):
                cell.number_format = "@"
        sheet.cell(row, 5).number_format = "0.000000"
        sheet.cell(row, 6).number_format = "0.000000"

    add_list_validation(
        sheet,
        "=year_list",
        f"C{FIRST_DATA_ROW}:C{final_row}",
        "Pilih tahun pembangunan antara 1900 dan tahun berjalan.",
    )

    add_custom_validation(
        sheet,
        "=AND(ISTEXT(A5),LEN(A5)>0,LEN(A5)<=50)",
        f"A{FIRST_DATA_ROW}:A{final_row}",
        "Kode/ID tidak valid",
        "Masukkan Kode/ID Menara sebagai teks maksimal 50 karakter.",
    )
    for column in ("B", "D", "G", "H", "J", "K", "L"):
        add_custom_validation(
            sheet,
            f"=AND(ISTEXT({column}5),LEN({column}5)>0,LEN({column}5)<=50)",
            f"{column}{FIRST_DATA_ROW}:{column}{final_row}",
            "Teks tidak valid",
            "Kolom wajib diisi dengan teks maksimal 50 karakter.",
        )

    add_custom_validation(
        sheet,
        "=AND(ISNUMBER(E5),E5>=95,E5<=141,ROUND(E5,6)=E5)",
        f"E{FIRST_DATA_ROW}:E{final_row}",
        "Longitude tidak valid",
        "Masukkan angka 95 sampai 141 dengan maksimal 6 angka desimal.",
    )
    add_custom_validation(
        sheet,
        "=AND(ISNUMBER(F5),F5>=-11,F5<=6,ROUND(F5,6)=F5)",
        f"F{FIRST_DATA_ROW}:F{final_row}",
        "Latitude tidak valid",
        "Masukkan angka -11 sampai 6 dengan maksimal 6 angka desimal.",
    )
    height_validation = DataValidation(
        type="whole",
        operator="between",
        formula1="1",
        formula2="999",
        allow_blank=False,
        errorStyle="stop",
        showErrorMessage=True,
        errorTitle="Tinggi menara tidak valid",
        error="Masukkan bilangan bulat dari 1 sampai 999.",
        showInputMessage=True,
        promptTitle="Tinggi menara",
        prompt="Bilangan bulat saja; rentang 1 sampai 999 meter.",
    )
    sheet.add_data_validation(height_validation)
    height_validation.add(f"I{FIRST_DATA_ROW}:I{final_row}")

    set_auto_width(sheet, minimum=14, maximum=42)
    for column, width in {
        "A": 27,
        "B": 36,
        "C": 34,
        "D": 36,
        "E": 17,
        "F": 17,
        "G": 38,
        "H": 35,
        "I": 36,
        "J": 40,
        "K": 42,
        "L": 44,
    }.items():
        sheet.column_dimensions[column].width = width
    return sheet


def create_year_master(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Data_Master")
    sheet.sheet_state = "hidden"
    sheet["A1"] = "Tahun Pembangunan"
    for row, year in enumerate(range(1900, 10000), 2):
        sheet.cell(row, 1, year)
    add_defined_name(
        workbook,
        "year_list",
        "'Data_Master'!$A$2:INDEX('Data_Master'!$A:$A,YEAR(TODAY())-1898)",
    )


def build_workbook(output_path: Path) -> None:
    workbook = Workbook()
    workbook.properties.title = "Template Registrasi Menara Telekomunikasi"
    workbook.properties.subject = "Template upload bulk registrasi menara"
    workbook.properties.creator = "Portal Infrastruktur Pasif"

    create_form_sheet(workbook)
    create_year_master(workbook)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="Output .xlsx path.",
    )
    args = parser.parse_args()
    build_workbook(args.output)
    print(f"Created: {args.output}")


if __name__ == "__main__":
    main()
